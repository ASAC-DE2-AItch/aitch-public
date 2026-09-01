# -*- coding: utf-8 -*-
"""
KB 5컬렉션 통합 개요 xlsx 생성 — 컬렉션당 1시트 + 요약 시트
================================================================================
목적 : rag_kb/<컬렉션>/*.jsonl 시드를 사람이 표로 훑어볼 수 있게 단일 xlsx 로 통합.
       (CSV 는 시트/탭이 없어 컬렉션별 분리가 안 됨 → xlsx 워크북 채택)
소스 : 각 컬렉션 폴더의 *_seed.jsonl / *.jsonl  (ingest 소스와 동일)
계약 : config/kb_schema.yaml — 공통 payload 14필드를 앞쪽에 고정 정렬.
사용 : python -m src.agent_service.export_kb_overview  [-o kb_5collections_overview.xlsx]
"""
import argparse
import json
import logging
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from src.agent_service.vectordb import RAG_KB_DIR, load_schema

log = logging.getLogger("kb-export")
logging.basicConfig(level=logging.INFO, format="%(levelname)s │ %(message)s")

HEADER_FILL = PatternFill("solid", fgColor="1F4E78")
HEADER_FONT = Font(color="FFFFFF", bold=True)
TITLE_FONT = Font(bold=True, size=13)
MAXW = 60   # 열 폭 상한(문자)


def _cell(v):
    """리스트/딕트는 JSON 문자열로, None 은 빈칸으로 평면화한다."""
    if v is None:
        return ""
    if isinstance(v, (list, dict)):
        return json.dumps(v, ensure_ascii=False)
    return v


def _load(logical: str) -> list:
    """컬렉션 폴더의 *.jsonl 전부를 레코드 리스트로 로드한다."""
    d = RAG_KB_DIR / logical
    recs = []
    for fp in sorted(d.glob("*.jsonl")):
        for line in fp.open(encoding="utf-8"):
            if line.strip():
                recs.append(json.loads(line))
    return recs


def _columns(logical: str, recs: list) -> list:
    """열 순서: 공통 14 → 전용 payload → 그 외 추가필드(등장순)."""
    schema = load_schema()
    common = list(schema.get("common_payload", {}))
    coll_payload = list((schema["collections"].get(logical, {}).get("payload") or {}))
    ordered, seen = [], set()
    for k in common + coll_payload:
        if k not in seen:
            ordered.append(k); seen.add(k)
    for r in recs:                       # 스키마에 없는 추가필드는 뒤에 등장순
        for k in r:
            if k not in seen:
                ordered.append(k); seen.add(k)
    return ordered


def _autosize(ws, ncols):
    """열 폭을 내용 길이에 맞춰 조정(상한 MAXW). 헤더행 고정·자동필터."""
    for c in range(1, ncols + 1):
        letter = get_column_letter(c)
        longest = max((len(str(ws.cell(r, c).value or "")) for r in range(1, ws.max_row + 1)), default=10)
        ws.column_dimensions[letter].width = min(max(longest + 2, 10), MAXW)
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(ncols)}{ws.max_row}"


def _write_collection(ws, logical: str, recs: list):
    """컬렉션 1개를 시트에 표로 기록한다."""
    cols = _columns(logical, recs)
    ws.append(cols)
    for c in range(1, len(cols) + 1):
        cell = ws.cell(1, c)
        cell.fill = HEADER_FILL; cell.font = HEADER_FONT
        cell.alignment = Alignment(vertical="center", wrap_text=True)
    for r in recs:
        ws.append([_cell(r.get(k)) for k in cols])
    _autosize(ws, len(cols))
    log.info(f"{logical}: {len(recs)}행 × {len(cols)}열")


def _write_summary(ws, rows):
    """요약 시트 — 컬렉션·건수·kind·목적·Postgres."""
    schema = load_schema()
    ws["A1"] = "KB 5컬렉션 개요"; ws["A1"].font = TITLE_FONT
    hdr = ["컬렉션(kb_type)", "건수", "kind", "목적", "Postgres 대응", "시트"]
    ws.append([]); ws.append(hdr)
    hr = ws.max_row
    for c in range(1, len(hdr) + 1):
        cell = ws.cell(hr, c); cell.fill = HEADER_FILL; cell.font = HEADER_FONT
    for logical, n in rows:
        cs = schema["collections"].get(logical, {})
        ws.append([logical, n, cs.get("kind", ""), cs.get("purpose", ""),
                   cs.get("postgres", ""), logical])
    ws.append([]); ws.append(["합계", sum(n for _, n in rows)])
    _autosize(ws, len(hdr)); ws.freeze_panes = "A1"


def build(out: Path):
    """전체 컬렉션을 읽어 요약 + 컬렉션별 시트로 xlsx 를 생성한다."""
    schema = load_schema()
    collections = list(schema["collections"])
    wb = Workbook()
    summary = wb.active; summary.title = "요약"
    rows = []
    for logical in collections:
        recs = _load(logical)
        rows.append((logical, len(recs)))
        _write_collection(wb.create_sheet(title=logical[:31]), logical, recs)
    _write_summary(summary, rows)
    wb.save(out)
    log.info(f"저장 완료: {out}  (시트 {len(collections)+1}개, 총 {sum(n for _, n in rows)}건)")


if __name__ == "__main__":
    ap = argparse.ArgumentParser()
    ap.add_argument("-o", default="kb_5collections_overview.xlsx")
    args = ap.parse_args()
    build(Path(args.o))
